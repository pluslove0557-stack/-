"""엑셀 마진 계산기 - 판매 마진/마진율/손익분기 판매가 계산 로직."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = [
    "상품명", "원가", "판매가", "배송비", "수수료율(%)",
    "광고비", "기타비용", "수량", "개당마진", "마진율(%)", "총마진",
]


def calc_row(row):
    """단일 상품 행에 대한 마진을 계산해서 dict로 반환한다."""
    cost = float(row.get("cost") or 0)
    price = float(row.get("price") or 0)
    shipping = float(row.get("shipping") or 0)
    fee_pct = float(row.get("fee_pct") or 0)
    ad_cost = float(row.get("ad_cost") or 0)
    misc_cost = float(row.get("misc_cost") or 0)
    qty = float(row.get("qty") or 1)

    fee_amount = price * fee_pct / 100
    unit_margin = price - cost - shipping - fee_amount - ad_cost - misc_cost
    margin_pct = (unit_margin / price * 100) if price else 0.0
    total_margin = unit_margin * qty

    denom = 1 - fee_pct / 100
    fixed_cost = cost + shipping + ad_cost + misc_cost
    breakeven_price = (fixed_cost / denom) if denom > 0 else None

    return {
        "name": row.get("name") or "",
        "cost": cost,
        "price": price,
        "shipping": shipping,
        "fee_pct": fee_pct,
        "ad_cost": ad_cost,
        "misc_cost": misc_cost,
        "qty": qty,
        "unit_margin": round(unit_margin, 2),
        "margin_pct": round(margin_pct, 2),
        "total_margin": round(total_margin, 2),
        "breakeven_price": round(breakeven_price, 2) if breakeven_price is not None else None,
    }


def calc_rows(rows):
    return [calc_row(r) for r in rows]


def build_workbook(rows):
    """입력값과 '엑셀 수식'이 그대로 살아있는 xlsx 워크북을 생성한다.

    다운로드 후에도 엑셀에서 값(원가/판매가 등)을 바꾸면 마진이 자동으로
    재계산되도록 수식으로 작성한다.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "마진계산"

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    start_row = 2
    for i, row in enumerate(rows):
        r = start_row + i
        ws.cell(row=r, column=1, value=row.get("name") or "")
        ws.cell(row=r, column=2, value=float(row.get("cost") or 0))
        ws.cell(row=r, column=3, value=float(row.get("price") or 0))
        ws.cell(row=r, column=4, value=float(row.get("shipping") or 0))
        ws.cell(row=r, column=5, value=float(row.get("fee_pct") or 0))
        ws.cell(row=r, column=6, value=float(row.get("ad_cost") or 0))
        ws.cell(row=r, column=7, value=float(row.get("misc_cost") or 0))
        ws.cell(row=r, column=8, value=float(row.get("qty") or 1))
        # 개당마진 = 판매가 - 원가 - 배송비 - (판매가*수수료율/100) - 광고비 - 기타비용
        ws.cell(row=r, column=9, value=f"=C{r}-B{r}-D{r}-(C{r}*E{r}/100)-F{r}-G{r}")
        # 마진율 = 개당마진 / 판매가 * 100
        ws.cell(row=r, column=10, value=f"=IF(C{r}=0,0,I{r}/C{r}*100)")
        # 총마진 = 개당마진 * 수량
        ws.cell(row=r, column=11, value=f"=I{r}*H{r}")

    if rows:
        total_row = start_row + len(rows)
        ws.cell(row=total_row, column=1, value="합계").font = Font(bold=True)
        ws.cell(row=total_row, column=11, value=f"=SUM(K{start_row}:K{total_row-1})").font = Font(bold=True)

    for col_idx in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    ws.column_dimensions["A"].width = 22
    ws.freeze_panes = "A2"

    return wb
