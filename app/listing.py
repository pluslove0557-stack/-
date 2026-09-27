"""자동 상품등록 프로그램 - 상품 리스트를 오픈마켓 대량등록 양식에 맞게 변환.

주의: 쿠팡/네이버 스마트스토어/11번가 등은 대량등록 엑셀 양식을 주기적으로
갱신하고, 실제 등록은 각 사의 판매자센터(또는 공식 오픈API)를 통해서만
가능하다. 이 도구는 "내가 가진 상품 데이터를 각 채널 업로드 양식의 컬럼에
맞춰 자동으로 정리"해 주는 매핑/변환 도구이며, 채널에 직접 자동으로
게시(크롤링/우회 등록)하지 않는다. 사용 전 반드시 각 채널 판매자센터에서
최신 공식 양식을 내려받아 컬럼을 비교/확인할 것.
"""

import csv
import io
from openpyxl import Workbook, load_workbook

# 일반적으로 오픈마켓 대량등록 양식에서 공통적으로 요구하는 항목들을
# 참고용 "표준 컬럼 세트"로 제공한다. 실제 채널별 정확한 컬럼명/순서는
# 반드시 각 판매자센터의 최신 양식으로 대체해서 사용해야 한다.
TEMPLATES = {
    "generic": {
        "label": "공통 표준 양식 (범용)",
        "notes": "특정 채널 지정 없이 상품 데이터를 정리할 때 사용하는 기본 컬럼 세트입니다.",
        "columns": [
            "상품명", "카테고리", "판매가", "원가", "재고수량",
            "옵션명", "옵션값", "대표이미지경로", "상세이미지경로",
            "배송비", "브랜드", "제조사", "바코드", "상품상태", "메모",
        ],
    },
    "coupang": {
        "label": "쿠팡(Wing) 대량등록 참고 양식",
        "notes": (
            "실제 등록은 쿠팡 Wing 판매자센터에서 다운로드한 최신 엑셀 양식을 사용하거나, "
            "쿠팡 Open API(상품 생성 API)로 진행해야 합니다. 아래 컬럼은 참고용입니다."
        ),
        "columns": [
            "등록상품명", "노출카테고리", "판매가격", "재고수량",
            "옵션명", "옵션값", "대표이미지URL", "추가이미지URL",
            "출고지", "반품지", "배송비종류", "배송비", "상품상태", "바코드",
        ],
    },
    "naver": {
        "label": "네이버 스마트스토어 대량등록 참고 양식",
        "notes": (
            "실제 등록은 네이버 커머스 API센터(스마트스토어센터 연동) 또는 "
            "스마트스토어센터에서 받은 최신 엑셀 양식을 사용해야 합니다."
        ),
        "columns": [
            "상품명", "카테고리ID", "판매가", "재고수량",
            "옵션유형", "옵션명", "옵션값", "대표이미지", "추가이미지",
            "배송비유형", "기본배송비", "원산지", "모델명", "제조사",
        ],
    },
    "st11": {
        "label": "11번가 대량등록 참고 양식",
        "notes": "실제 등록은 11번가 셀러오피스에서 받은 최신 엑셀 양식 또는 오픈API를 사용해야 합니다.",
        "columns": [
            "상품명", "카테고리코드", "판매가", "재고수량",
            "옵션명1", "옵션값1", "대표이미지", "상세이미지",
            "배송비구분", "배송비", "원산지", "브랜드",
        ],
    },
}


def list_templates():
    return {k: {"label": v["label"], "notes": v["notes"]} for k, v in TEMPLATES.items()}


def get_template(key):
    return TEMPLATES.get(key)


def read_source_table(filename, file_bytes):
    """업로드된 CSV/XLSX 파일을 (headers, rows[list[dict]]) 로 파싱한다."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h) if h is not None else "" for h in next(rows_iter)]
        rows = []
        for raw in rows_iter:
            if raw is None or all(v is None for v in raw):
                continue
            rows.append({headers[i]: raw[i] for i in range(len(headers)) if i < len(raw)})
        return headers, rows

    text = file_bytes.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    headers = reader.fieldnames or []
    rows = [dict(r) for r in reader]
    return headers, rows


def apply_mapping(source_rows, target_columns, mapping):
    """mapping: {target_column: source_column_or_literal}.

    source_column 이 실제 소스 헤더에 존재하면 그 값을 채우고,
    존재하지 않으면 고정값(리터럴)로 간주해 모든 행에 그대로 채운다.
    비어 있으면 빈 문자열로 채운다.
    """
    out_rows = []
    for src in source_rows:
        out = {}
        for col in target_columns:
            key = mapping.get(col, "")
            if not key:
                out[col] = ""
            elif key in src:
                out[col] = src.get(key, "")
            else:
                out[col] = key  # literal / 고정값
        out_rows.append(out)
    return out_rows


def build_output_workbook(target_columns, out_rows, sheet_title="상품등록"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(target_columns)
    for row in out_rows:
        ws.append([row.get(c, "") for c in target_columns])
    for i, _ in enumerate(target_columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 18
    ws.freeze_panes = "A2"
    return wb
